"""Construye la base de datos auditable en Excel: 03_Base_de_Datos/RECOEVO_Base_de_Datos_Campo.xlsx"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE = Path(__file__).resolve().parents[1]
DB = BASE / "03_Base_de_Datos"
CB = json.loads((BASE / "02_Metodologia" / "codebook.json").read_text(encoding="utf-8"))
R = json.loads((DB / "resultados.json").read_text(encoding="utf-8"))
U = json.loads((DB / "unidades_codificadas.json").read_text(encoding="utf-8"))
V = json.loads((BASE / "01_Validacion_Transcripciones" / "validacion.json").read_text(encoding="utf-8"))
FF = {u["id"]: u for u in json.loads((DB / "frases_finales.json").read_text(encoding="utf-8"))}
mmss = lambda t: "" if t is None else f"{int(t // 60):02d}:{int(t % 60):02d}"
NAME = {c["id"]: c["nombre"] for c in CB["codigos"]}
DIMN = {d["id"]: d["nombre"] for d in CB["dimensiones"]}
CARA = {1: "😠", 2: "🙁", 3: "😐", 4: "🙂", 5: "😄"}
SAT = {1: "Muy insatisfecho", 2: "Insatisfecho", 3: "Neutral", 4: "Satisfecho", 5: "Muy satisfecho"}
DIS = {1: "Rechazo", 2: "Escéptico", 3: "Condicionado", 4: "Favorable", 5: "Entusiasta"}

HEAD = PatternFill("solid", fgColor="0B1014"); HF = Font(bold=True, color="B8F35A", name="Calibri", size=11)
thin = Border(bottom=Side(style="thin", color="D0D5D2"))
wb = Workbook()

def sheet(title, headers, rows, widths, table=True, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for r in rows: ws.append(r)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]: c.fill = HEAD; c.font = HF; c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    if table and rows:
        t = Table(displayName=title.replace(" ", "_").replace("·", "")[:30], ref=f"A1:{get_column_letter(len(headers))}{len(rows)+1}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True); ws.add_table(t)
    return ws

# 1 · Resumen
esc = R["escalas"]; ac = R["acuerdo"]
res = [
    ["Participantes únicos", 8, "3 hogares, 3 comercios, 2 hogar y comercio"],
    ["Grabaciones", 10, "Audio2+3 = E02; Audio6+10 = E05 (probable)"],
    ["Minutos de audio", round(sum(v["duracion_s"] for v in V) / 60, 1), "Suma de duraciones de archivo"],
    ["Unidades de significado", R["n_unidades"], "Turnos del participante divididos en frases ≤55 palabras"],
    ["Unidades con contenido", R["n_sustantivas"], "Con al menos un código"],
    ["Códigos", len(CB["codigos"]), "5 dimensiones · deductivos + inductivos"],
    ["Satisfacción media (1-5)", esc["satisfaccion"]["media"], f"n = {esc['satisfaccion']['n']} frases que evalúan el barrio o el servicio · 7 de 8 personas con promedio < 2,5"],
    ["Disposición media (1-5)", esc["disposicion"]["media"], f"n = {esc['disposicion']['n']} opiniones sobre soluciones"],
    ["Kappa de Cohen global (A vs B)", ac["kappa_global"], f"Acuerdo bruto {ac['acuerdo_global']:.1%} · 220 × 17 decisiones"],
    ["Kappa ponderado satisfacción", ac["kappa_pond_satisfaccion"], "Ponderación cuadrática"],
    ["Kappa ponderado disposición", ac["kappa_pond_disposicion"], "Ponderación cuadrática"],
    ["Entrevistas para 80 % de temas", R["entrevistas_para_80pct"], "Promedio de 1.000 órdenes aleatorios"],
    ["Temas con 4 entrevistas", f"{R['saturacion_permutada'][3]} de 17", "Saturación de códigos (97 %)"],
]
ws = sheet("Resumen", ["Indicador", "Valor", "Nota"], res, [36, 16, 70], first=True)

# 2 · Frases
rows = []
for u in U:
    f = FF.get(u["id"], {})
    rows.append([u["id"], u["audio"], mmss(f.get("t")), u["participante"], R["participantes"][u["participante"]]["rol"], u["pregunta"], f.get("e", ""), u["texto"], u["palabras"],
                 ", ".join(u["codigos"]), " | ".join(NAME[c] for c in u["codigos"]), ", ".join(DIMN[d] for d in u["dimensiones"]),
                 u["satisfaccion"], CARA.get(u["satisfaccion"], ""), SAT.get(u["satisfaccion"], ""),
                 u["disposicion"], CARA.get(u["disposicion"], ""), DIS.get(u["disposicion"], ""),
                 ", ".join(u["codigos_B"]), u["satisfaccion_B"], u["disposicion_B"],
                 "Sí" if set(u["codigos"]) == set(u["codigos_B"]) else "No", ", ".join(u["condiciones"]), u["nota"]])
ws = sheet("Frases codificadas", ["ID", "Audio", "Minuto", "Participante", "Rol", "Pregunta previa", "Frase editada (gramática)", "Frase literal (transcripción)", "Palabras",
            "Códigos A", "Temas", "Dimensión", "Satisfacción", "Carita", "Nivel satisfacción", "Disposición", "Carita disp.", "Nivel disposición",
            "Códigos B (Qwen 14B)", "Satisf. B", "Disp. B", "Acuerdo exacto A=B", "Condiciones para participar", "Nota del codificador"], rows,
           [12, 9, 8, 11, 18, 40, 60, 60, 8, 14, 32, 24, 11, 7, 16, 11, 7, 14, 16, 9, 9, 10, 20, 30])
for col in ("M", "P"):
    ws.conditional_formatting.add(f"{col}2:{col}{len(rows)+1}", ColorScaleRule(start_type="num", start_value=1, start_color="FF5A5A", mid_type="num", mid_value=3, mid_color="FFD24D", end_type="num", end_value=5, end_color="4CD98A"))
for r in range(2, len(rows) + 2):
    for col in ("N", "Q"): ws[f"{col}{r}"].font = Font(size=16); ws[f"{col}{r}"].alignment = Alignment(horizontal="center", vertical="top")

# 3 · Libro de códigos
sheet("Libro de codigos", ["ID", "Dimensión", "Código", "Origen", "Definición", "Incluye", "Excluye", "Participantes (de 8)", "Frases", "Satisfacción media", "Kappa A vs B"],
      [[c["id"], DIMN[c["dim"]], c["nombre"], c["origen"], c["definicion"], c["incluye"], c["excluye"], R["frecuencias"][c["id"]]["n_part"],
        R["frecuencias"][c["id"]]["unidades"], esc["satisfaccion_por_codigo"][c["id"]], ac["por_codigo"].get(c["id"], {}).get("kappa")] for c in CB["codigos"]],
      [7, 26, 28, 11, 60, 40, 30, 13, 9, 12, 11])

# 4 · Matriz participante × código
codes = [c["id"] for c in CB["codigos"]]
mrows = [[p, R["participantes"][p]["rol"]] + [R["matriz"][p][c] for c in codes] for p in sorted(R["matriz"])]
mrows.append(["Personas", ""] + [R["frecuencias"][c]["n_part"] for c in codes])
ws = sheet("Matriz caso x codigo", ["Participante", "Rol"] + [NAME[c] for c in codes], mrows, [12, 20] + [11] * len(codes), table=False)
ws.row_dimensions[1].height = 70
ws.conditional_formatting.add(f"C2:{get_column_letter(len(codes)+2)}9", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="max", end_color="4CD98A"))

# 5 · Co-ocurrencias
sheet("Coocurrencias", ["Código A", "Código B", "Frases donde aparecen juntos"], [[NAME[c["a"]], NAME[c["b"]], c["n"]] for c in R["coocurrencias"]], [30, 30, 16])

# 6 · Escalas
sheet("Escalas", ["Participante", "Rol", "Satisfacción media", "Carita", "Disposición media", "Carita disp."],
      [[p, R["participantes"][p]["rol"], v["satisfaccion"], CARA.get(round(v["satisfaccion"])) if v["satisfaccion"] else "",
        v["disposicion"], CARA.get(round(v["disposicion"])) if v["disposicion"] else "—"] for p, v in esc["por_participante"].items()],
      [12, 22, 16, 8, 16, 10])

# 7 · Saturación
sheet("Saturacion", ["Entrevistas acumuladas", "Temas distintos (promedio 1.000 órdenes)", "% de 17"],
      [[i + 1, v, round(v / 17, 3)] for i, v in enumerate(R["saturacion_permutada"])], [22, 36, 10])

# 8 · Validación transcripciones
sheet("Validacion transcripcion", ["Audio", "Duración (s)", "Palabras audio (Whisper)", "Palabras transcripción", "Cobertura", "WER aprox.", "Similitud", "Decisión"],
      [[v["audio"], v["duracion_s"], v["palabras_audio"], v["palabras_transcripcion"], v["cobertura"], v["wer_aprox"], v["similitud"],
        {"Audio6": "Validada tras escucha dirigida de bloques dudosos", "Audio7": "Reemplazada por versión revisada (hablantes asignados)",
         "Audio10": "Reemplazada por versión revisada; solo 28 s sustantivos"}.get(v["audio"], "Validada")] for v in V],
      [10, 12, 18, 18, 11, 11, 11, 50])

# 9 · Evidencia visual
EV = [
    ("WhatsApp Image … 10.44.51 AM (1)/(2)", "Entrevista en la calle (equipo)", "Método", "—", "No (duplicada)"),
    ("WhatsApp Image … 10.44.51 AM (3)", "Mural “Cuida tu shut” junto a ladera con residuos y perro", "PUN, FAU", "Dicho vs visto: ironía del mural", "Sí · L07"),
    ("WhatsApp Image … 10.44.51 AM", "Bolsas grises sobre el andén", "PUN", "“Una noche antes sacan la basura”", "Sí · L19"),
    ("WhatsApp Image … 10.44.52 AM (1)", "Bolsas cerradas en la acera", "PUN", "Bolsa expuesta", "No"),
    ("WhatsApp Image … 10.44.52 AM", "Perro negro entre residuos en ladera", "FAU, ROM", "“Los perros rompen las bolsas”", "Sí · L19"),
    ("WhatsApp Image … 10.44.53 AM", "Tienda con vista al shut rosado y basural", "PUN", "El shut como paisaje cotidiano", "Sí · L07"),
    ("WhatsApp Image … 10.44.58 AM / 10.45.00 AM", "Perros frente al shut rosado desbordado", "FAU, PUN", "“Hay más perros que gente”", "Sí · L09"),
    ("WhatsApp Image … 10.45.01 AM", "Shut con cartel “¡Prohibido!” desbordado, perro dentro", "PUN, EST", "Norma sin control", "Sí · L06, L07"),
    ("WhatsApp Image … 10.45.02 AM (1)", "Llanta con cáscaras de banano y plásticos mezclados", "MEZ", "“Todo es un revoltijo”", "Sí · L19"),
    ("WhatsApp Image … 10.45.02 AM (2)", "Caneca de pedal con bolsa en un local", "ART", "El pedal ya es un gesto familiar", "Referida · L22"),
    ("WhatsApp Image … 10.45.02 AM", "Perro echado en pasto con papeles", "FAU", "Perros en calle", "No"),
    ("WhatsApp Image … 11.12.54 AM", "Fruver con clientes (aparece una menor)", "APR", "Fuente de orgánicos", "No (menor de edad)"),
    ("WhatsApp Image … 11.12.59 AM / 11.13.00 AM", "Entrevista en fruver", "Método", "—", "No"),
    ("WhatsApp Image … 11.13.00 AM (1)/(2)", "Firma del consentimiento sobre la báscula", "Ética", "Autorización", "No"),
    ("WhatsApp Image … 11.13.01 AM (1)", "Monumento / aviso de bienvenida al barrio", "COM", "E03 menciona el aviso", "No"),
    ("WhatsApp Image … 11.13.01 AM", "Vía destapada con charcos frente a viviendas", "AMB", "“Cuando llueve se forma un pichal”", "Sí · L19"),
    ("WhatsApp Image … 7.47.18 AM (1)/(2)", "Punto de arrojo junto a arbusto; persona y palomas", "PUN, FAU", "Acopio informal", "No (persona)"),
    ("WhatsApp Video … 7.47.02 AM (13 s)", "Personas llevan bolsas y cartón al botadero; palomas y perro", "PUN, FAU", "Disposición en espacio público", "Sí · L01 portada"),
    ("WhatsApp Video … 7.47.11 AM (11 s)", "Recicladores manipulan residuos con perros alrededor", "ROM, FAU", "“Rompen las bolsas”", "Sí · L09 (desde el s 5, sin rostro)"),
    ("WhatsApp Video … 7.47.32 AM (10 s)", "Vendedor ambulante de fruta, cáscaras en el suelo, perro", "APR, PUN", "Generación de orgánicos en la calle", "Sí · L25 (difuminado)"),
    ("Basuras.png (Google Street View)", "Calle en pendiente, parte alta del barrio: residuos acumulados en una esquina, aparte del shut", "PUN, AMB", "El problema se extiende a varias esquinas; pendientes", "Sí · O1 (con fuente)"),
    ("WhatsApp Video … 11.12.59 AM (18 s)", "Interior del fruver, cámara en movimiento", "Método", "—", "No (plano inestable)"),
]
sheet("Evidencia visual", ["Archivo", "Qué se ve (solo lo visible)", "Códigos vinculados", "Relación con lo dicho", "¿En la presentación?"], [list(e) for e in EV], [40, 55, 16, 36, 26])

# 10 · Participantes
sheet("Participantes", ["Código", "Rol", "Perfil (anonimizado)", "Grabaciones"],
      [[p, v["rol"], v["detalle"], " + ".join(v["audios"])] for p, v in R["participantes"].items()], [10, 22, 55, 18])

out = DB / "RECOEVO_Base_de_Datos_Campo.xlsx"
wb.save(out); print("ok", out)
